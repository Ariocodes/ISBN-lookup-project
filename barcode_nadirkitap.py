#   NADIRKITAP.COM SCRAPER PROJECT
#   DEVELOPED BY Ario Bashiri
#   MAIL: ario@ariobashiri.com
#   MAIL: ariobashiri@gmail.com
#   GITHUB: https://github.com/ariocodes


import cv2
import re
import sys
from pyzbar.pyzbar import decode
from datetime import datetime
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
import cloudscraper
import tkinter as tk
from PIL import Image, ImageTk


# CONFIGS
OUTPUT_FILE = "nadirkitap_results.xlsx"
BASE_URL = "https://www.nadirkitap.com/kitapara.php?ara=aramayap&ref=anasayfa&tip=kitap&isbn={isbn}"
scraper = cloudscraper.create_scraper()



# EXCEL COLUMNS (These are the data that will be scraped :P )
COLUMNS = [
    "ISBN",
    "Book Name",
    "Author",
    "Number of Listings",
    "Prices",
    "Average Price",
    "Min Price",
    "Max Price",
    "1★ avg",
    "2★ avg",
    "3★ avg",
    "4★ avg",
    "5★ avg",
    "\"Yeni\" Avg",
    "URL",
    "Scrape Date"
]



# THE SCRAPER
def scrape_isbn(isbn: str):
    url = BASE_URL.format(isbn=isbn)
    print(url)
    try:
        resp = scraper.get(url, timeout=15)
    except Exception as e:
        print("Request failed:", e)
        return []

    if resp.status_code != 200:
        print(f"Blocked or error: {resp.status_code}")
        return []

    soup = BeautifulSoup(resp.text, "html.parser")
    results = []
    cards = soup.find_all("div")
    for card in cards:
        # finding title
        title_tag = card.find("h4")
        if not title_tag:
            continue
        # finding book name
        a_tag = title_tag.find("a")
        if not a_tag:
            continue

        book_name = a_tag.get_text(strip=True)

        # saving the ad's link
        link = a_tag.get("href", "")
        if link and not link.startswith("http"):
            link = "https://www.nadirkitap.com" + link

        text_block = card.get_text(" ", strip=True)

        price_match = re.search(r"\d+[.,]\d+\s*TL", text_block)
        price = price_match.group(0) if price_match else ""

        author = ""
        tag = card.select_one("span.nk-author-text")

        if tag:
            author = tag.get_text(strip=True)
        else:
            spans = card.select("span.media-tooltip")
            for s in spans:
                t = s.get_text(strip=True)
                if "Yayınları" not in t and re.search(r"[A-Za-zÇĞİÖŞÜ]", t):
                    author = t
                    break

        condition = ""
        cond_tag = card.select_one("span.text-nadir")
        if cond_tag:
            condition = cond_tag.get_text(" ", strip=True)
            condition = re.sub(r"\s+", " ", condition).strip()

        results.append({
            "ISBN": isbn,
            "Book Name": book_name,
            "Author": author,
            "Price": price,
            "Condition": condition,
            "Listing URL": link,
            "Source URL": url,
            "Scraped At": datetime.now().strftime("%Y-%m-%d %H:%M")
        })
    return results


# EXPORTING TO AN EXCEL
def save_excel(records):
    if not records:
        return

    try:
        wb = load_workbook(OUTPUT_FILE)
        ws = wb.active
    except:
        wb = Workbook()
        ws = wb.active
        ws.append(COLUMNS)

    for r in records:
        ws.append([r.get(c, "") for c in COLUMNS])

    wb.save(OUTPUT_FILE)
    print("Saved to Excel")




# GLOBAL STATES
seen = set()
records = []
cap = cv2.VideoCapture(0)





# PROCESSING ISBN
def process_isbn(isbn, source_url):
    if isbn in seen:
        return
    seen.add(isbn)

    print("\nDetected ISBN:", isbn)

    listings = scrape_isbn(isbn)
    if not listings:
        print("No listings found")
        return

    prices = []
    cond_map = {1: [], 2: [], 3: [], 4: [], 5: []}
    yeni_prices = []
    authors = set()

    for item in listings:

        author = item.get("Author")
        if not author:
            continue
        authors.add(author.strip())

        price_text = item.get("Price", "")
        match = re.search(r"\d+[.,]\d+", price_text)

        price = None
        if match:
            price = float(match.group().replace(",", "."))
            prices.append(price)

        condition_text = item.get("Condition", "")
        text = condition_text.lower().strip()

        # STRICT "YENI ONLY" CHECK
        # must contain ONLY the word "yeni" (no extra words)
        is_pure_yeni = (text == "yeni")

        # STAR COUNT
        star = condition_text.count("★")
        if star < 1 or star > 5:
            star = None

        if price is not None:

            # 1–5 star buckets
            if star:
                cond_map[star].append(price)

            # ONLY pure "Yeni"
            if is_pure_yeni:
                yeni_prices.append(price)


    avg = lambda lst: round(sum(lst) / len(lst), 2) if lst else "-"

    record = {
        "ISBN": isbn,
        "Book Name": listings[0].get("Book Name", ""),
        "Author": ", ".join(authors) if authors else "-",
        "Number of Listings": len(listings),
        "Prices": ", ".join(map(str, prices)) if prices else "-",
        "Average Price": avg(prices),
        "Min Price": min(prices) if prices else "-",
        "Max Price": max(prices) if prices else "-",

        "1★ avg": avg(cond_map[1]),
        "2★ avg": avg(cond_map[2]),
        "3★ avg": avg(cond_map[3]),
        "4★ avg": avg(cond_map[4]),
        "5★ avg": avg(cond_map[5]),
        "\"Yeni\" Avg": avg(yeni_prices),

        "URL": source_url,
        "Scrape Date": datetime.now().strftime("%Y-%m-%d %H:%M")
    }

    records.append(record)

    # KEEPING THE EXACT OUTPUT FORMAT
    print(
        f"Book Name: {record['Book Name']} | "
        f"Author: {record['Author']} | \n"
        f"Listings: {len(listings)} | "
        f"Avg Price: {record['Average Price']} | "
        f"1★:{record['1★ avg']} 2★:{record['2★ avg']} 3★:{record['3★ avg']} 4★:{record['4★ avg']} 5★:{record['5★ avg']} \"Yeni\":{record['\"Yeni\" Avg']} |"
    )


# TKINTER GUI
root = tk.Tk()
root.title("ISBN Scanner")

video_label = tk.Label(root)
video_label.pack()

entry = tk.Entry(root, font=("Arial", 14))
entry.bind("<Return>", lambda event: submit_isbn())

entry = tk.Entry(root, font=("Arial", 14))
entry.pack(pady=(15, 5))   # space above and below

entry.pack()

def submit_isbn():
    isbn = entry.get().strip()
    if isbn:
        process_isbn(isbn, BASE_URL.format(isbn=isbn))
        entry.delete(0, tk.END)

tk.Button(root, text="Add ISBN", command=submit_isbn).pack(pady=(5, 15))


# SAFE EXIT (FIXED SAVE)
def on_close():
    print("Closing app... saving Excel")
    save_excel(records)
    cap.release()
    root.destroy()

root.protocol("WM_DELETE_WINDOW", on_close)


# CAMERA LOOP + GREEN BOX FIX
def update_frame():
    ret, frame = cap.read()
    if ret:

        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        barcodes = decode(gray)

        for bc in barcodes:
            isbn = bc.data.decode("utf-8")

            bx, by, bw, bh = bc.rect

            # GREEN BOX RESTORED
            cv2.rectangle(frame, (bx, by), (bx + bw, by + bh), (0, 255, 0), 2)
            cv2.putText(frame, isbn, (bx, by - 10),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)

            process_isbn(isbn, BASE_URL.format(isbn=isbn))

        img = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        img = Image.fromarray(img)
        imgtk = ImageTk.PhotoImage(image=img)

        video_label.imgtk = imgtk
        video_label.configure(image=imgtk)

    root.after(10, update_frame)


# running each one 
update_frame()
root.mainloop()
cap.release()


#  /\_/\
# ( o.o )
#  > ^ <
# Mr. Meow

